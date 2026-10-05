"""Reusable python-docx toolkit used to build SmartHelp documentation .docx files.

Mirrors the visual language of the existing SmartPOS deliverable
(generate_video_script.py / SmartPOS-Documentation.docx): Calibri body, dark
navy headings, shaded code blocks, and 'Light * Accent 1' tables.
"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ── Palette (mirrors SmartPOS brand) ────────────────────────────────────────
NAVY = RGBColor(0x0B, 0x1B, 0x3A)
BLUE = RGBColor(0x1E, 0x5F, 0xE8)
GREEN = RGBColor(0x1B, 0xA3, 0x52)
RED = RGBColor(0xC0, 0x25, 0x25)
AMBER = RGBColor(0xB4, 0x6A, 0x0B)
GRAY = RGBColor(0x64, 0x74, 0x8B)
DARKGRAY = RGBColor(0x33, 0x3D, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CODE_BG = "F4F6F8"
INLINE_BG = "EEF2F7"
NOTE_BG = "F6F9FC"
WARN_BG = "FEF6E7"

_INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")


def new_document() -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0

    for level, size in ((1, 18), (2, 14), (3, 12), (4, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(11 if level <= 2 else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    _add_code_style(doc)
    _add_inline_style(doc)
    _add_note_style(doc)
    _add_caption_style(doc)

    section = doc.sections[0]
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)

    _build_running_header_footer(section)
    return doc


# ── Low-level style helpers ────────────────────────────────────────────────


def _add_code_style(doc: Document) -> None:
    style = doc.styles.add_style("SH Code", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(8)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.left_indent = Inches(0.06)
    style.paragraph_format.right_indent = Inches(0.04)
    p_pr = style.element.get_or_add_pPr()
    _shade(p_pr, CODE_BG)
    _borders(p_pr, color="C8D2DE", size=18, sides=("left",))
    # Consolas must be applied to ascii/hAnsi as well for Word on Windows.
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attr), "Consolas")


def _add_inline_style(doc: Document) -> None:
    style = doc.styles.add_style("SH Inline", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Consolas"
    style.font.size = Pt(9.5)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attr), "Consolas")
    _shade(r_pr, INLINE_BG)


def _add_note_style(doc: Document) -> None:
    style = doc.styles.add_style("SH Note", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.italic = True
    style.font.color.rgb = DARKGRAY
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.left_indent = Inches(0.12)
    _shade(style.element.get_or_add_pPr(), NOTE_BG)


def _add_caption_style(doc: Document) -> None:
    style = doc.styles.add_style("SH Caption", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(9)
    style.font.italic = True
    style.font.color.rgb = GRAY
    style.paragraph_format.space_before = Pt(2)
    style.paragraph_format.space_after = Pt(10)


def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _borders(p_pr, color: str = "D8DEE6", size: int = 4, sides=("left",)) -> None:
    p_bdr = OxmlElement("w:pBdr")
    for side in sides:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), color)
        p_bdr.append(node)
    p_pr.append(p_bdr)


def _build_running_header_footer(section) -> None:
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("SmartHelp  ·  Product & Technical Documentation")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = footer.add_run("SmartHelp — internal build specification  |  page ")
    left.font.size = Pt(8)
    left.font.color.rgb = GRAY
    _field(footer, "PAGE")
    tail = footer.add_run(" of ")
    tail.font.size = Pt(8)
    tail.font.color.rgb = GRAY
    _field(footer, "NUMPAGES")


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ── Inline rich text ───────────────────────────────────────────────────────


def add_runs(paragraph, text: str, size: float | None = None, color=None, bold=False):
    """Render a line of text supporting **bold**, *italic* and `code` spans."""
    for part in _INLINE_TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.style = paragraph.part.document.styles["SH Inline"]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.bold = bold or None
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


# ── Block helpers ──────────────────────────────────────────────────────────


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def h4(doc, text):
    return doc.add_heading(text, level=4)


def title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)
    if subtitle:
        sub = doc.add_paragraph()
        sub_run = sub.add_run(subtitle)
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = BLUE
        sub.paragraph_format.space_after = Pt(10)
    return p


def p(doc, text, size=None, color=None, space_after=6):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    add_runs(para, text, size=size, color=color)
    return para


def note(doc, text):
    para = doc.add_paragraph(style="SH Note")
    add_runs(para, text)
    return para


def bullets(doc, items, style="List Bullet", indent=0.18):
    for item in items:
        para = doc.add_paragraph(style=style)
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.left_indent = Inches(indent + 0.12)
        add_runs(para, item, size=10)


def steps(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Number")
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.left_indent = Inches(0.3)
        add_runs(para, item, size=10)


def code(doc, text, caption=None):
    """One paragraph, soft line breaks - a code block, not N stacked paragraphs."""
    lines = text.strip("\n").split("\n")
    para = doc.add_paragraph(style="SH Code")
    for index, line in enumerate(lines):
        if index:
            para.add_run().add_break(WD_BREAK.LINE)
        if line:
            para.add_run(line)
    if caption:
        cap(doc, caption)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return para


def diagram(doc, text, caption=None):
    lines = text.strip("\n").split("\n")
    para = doc.add_paragraph(style="SH Code")
    for index, line in enumerate(lines):
        if index:
            para.add_run().add_break(WD_BREAK.LINE)
        run = para.add_run(line)
        run.bold = index == 0 and line.strip().startswith("=")
    if caption:
        cap(doc, caption)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return para


def cap(doc, text):
    para = doc.add_paragraph(style="SH Caption")
    add_runs(para, text)
    return para


def callout(doc, label, text, fill=NOTE_BG, label_color=BLUE):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.rows[0].cells[0]
    _shade(cell._tc.get_or_add_tcPr(), fill)
    para = cell.paragraphs[0]
    run = para.add_run(f"{label}  ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = label_color
    add_runs(para, text, size=10)
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def table(doc, headers, rows, widths=None, style="Light Grid Accent 1", font_size=9):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = style
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    head = tbl.rows[0]
    _repeat_header(head)
    for index, header in enumerate(headers):
        cell = head.cells[index]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.space_before = Pt(1)
        run = para.add_run(str(header))
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = NAVY

    for row_values in rows:
        cells = tbl.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(1)
            add_runs(para, "" if value is None else str(value), size=font_size)

    if widths:
        for row in tbl.rows:
            for index, width in enumerate(widths):
                if index < len(row.cells):
                    row.cells[index].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def kv_table(doc, pairs, style="Light Shading Accent 1", label_width=1.85, value_width=4.6):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = style
    for label, value in pairs:
        cells = tbl.add_row().cells
        for index, text in enumerate((label, value)):
            cell = cells[index]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(1)
            run = para.add_run(text)
            run.font.size = Pt(9.5)
            if index == 0:
                run.bold = True
        cells[0].width = Inches(label_width)
        cells[1].width = Inches(value_width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def page_break(doc):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    para.add_run().add_break(WD_BREAK.PAGE)


def part_divider(doc, number, title_text, blurb):
    page_break(doc)
    p1 = doc.add_paragraph()
    run = p1.add_run(f"PART {number}")
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = BLUE
    p1.paragraph_format.space_after = Pt(0)
    p2 = doc.add_paragraph(style="Heading 1")
    run2 = p2.add_run(title_text)
    run2.font.size = Pt(24)
    run2.font.bold = True
    run2.font.color.rgb = NAVY
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    p3 = doc.add_paragraph()
    run3 = p3.add_run(blurb)
    run3.font.size = Pt(10.5)
    run3.font.color.rgb = GRAY
    p3.paragraph_format.space_after = Pt(12)
    rule = doc.add_paragraph()
    _borders(rule._p.get_or_add_pPr(), color="1E5FE8", size=12, sides=("bottom",))
    rule.paragraph_format.space_after = Pt(10)


def landscape_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    width, height = section.page_height, section.page_width
    section.page_width, section.page_height = width, height
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    _build_running_header_footer(section)
    return section


def portrait_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    _build_running_header_footer(section)
    return section


def toc_field(doc, levels: str = "1-3") -> None:
    para = doc.add_paragraph()
    run = para.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click this table of contents and choose “Update Field” to build it."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, placeholder, end):
        run._r.append(node)


def update_fields_on_open(doc) -> None:
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
